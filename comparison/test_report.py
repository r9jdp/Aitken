"""Verify report numbers, citation relationships and output scope without training."""
import csv
import json
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from build_report import HERE, ROOT, NAME, output_path, digest


class ReportEvidenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = Document(HERE/(NAME+'.docx'))
        cls.e = json.loads((HERE/'evidence/high_pollution_diagnosis.json').read_text())
        cls.cap = json.loads((ROOT/'results/outlier_sensitivity_20260909/results.json').read_text())

    def test_original_model_table(self):
        with (ROOT/'results/metrics.csv').open(newline='') as f:
            records = list(csv.DictReader(f))
        rows = self.doc.tables[0].rows[1:]
        self.assertEqual(len(rows), 7)
        for row, record in zip(rows, records):
            self.assertEqual([c.text for c in row.cells], [record['model']] + [f"{float(record[k]):.4f}" for k in ('mae_ug_m3','rmse_ug_m3','r2')])

    def test_error_conservation(self):
        self.assertEqual(self.e['below_25']['n']+self.e['at_least_25']['n'], 9619)
        self.assertAlmostEqual(self.e['below_25']['squared_error_percent']+self.e['at_least_25']['squared_error_percent'], 100)
        self.assertEqual(self.e['at_least_25']['underprediction_count'], 216)
        self.assertTrue(all(self.e['verification'].values()))
        rows = self.doc.tables[2].rows[1:]
        for row, key in zip(rows, ('overall','below_25','at_least_25')):
            d = self.e[key]
            self.assertEqual([c.text for c in row.cells][1:], [f"{d['n']:,}",f"{d['mae']:.2f}",f"{d['rmse']:.2f}",f"{d['bias']:.2f}",f"{d['squared_error_percent']:.2f}%"])

    def test_screening_matches_completed_experiment(self):
        rows = self.doc.tables[4].rows[1:]
        expected = [self.cap['screening'][m][c]['metrics'] for m in ('standalone','hybrid') for c in ('control','cap_995','cap_990')]
        for row, d in zip(rows, expected):
            self.assertEqual([c.text for c in row.cells][1:4], [f"{d['rmse']:.4f}",f"{d['mae']:.4f}",f"{d['high']['mae']:.4f}"])

    def test_matched_final_and_negative_verdict(self):
        rows = self.doc.tables[5].rows[1:]
        final = self.cap['final']['standalone']
        for row, key in zip(rows, ('control','cap_990')):
            d = final[key]
            self.assertEqual([c.text for c in row.cells][1:], [f"{d['mae']:.4f}",f"{d['rmse']:.4f}",f"{d['r2']:.4f}",f"{d['high']['mae']:.4f}"])
        self.assertGreater(final['paired_block_interval']['ci95'][0], 0)
        self.assertIn('Verdict: do not adopt capping.', (HERE/'REPORT.md').read_text(encoding='utf-8'))

    def test_inputs_unchanged(self):
        audit = json.loads((HERE/'evidence/build_audit.json').read_text())
        for relative, expected in audit['input_hashes_unchanged'].items():
            self.assertEqual(digest(ROOT/relative), expected)

    def test_outputs_cannot_escape_comparison(self):
        with self.assertRaises(ValueError):
            output_path('../outside.txt')
        self.assertTrue(output_path('README.md').is_relative_to(HERE))

    def test_all_sources_are_external_clickable_links(self):
        sources = json.loads((HERE/'sources.json').read_text(encoding='utf-8'))['sources']
        with ZipFile(HERE/(NAME+'.docx')) as z:
            rels = ET.fromstring(z.read('word/_rels/document.xml.rels'))
            links = {r.attrib['Target'] for r in rels if r.attrib.get('TargetMode') == 'External'}
            for source in sources:
                self.assertIn(source['url'], links)
            xml = ET.fromstring(z.read('word/document.xml'))
            w = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            self.assertEqual(len(xml.findall('.//w:br[@w:type="page"]',w)), 7)
            self.assertEqual(len(xml.findall('.//w:tbl',w)), 7)
            self.assertEqual(len(xml.findall('.//w:tblHeader',w)), 7)
            styles = ET.fromstring(z.read('word/styles.xml'))
            self.assertEqual(len(styles.findall('.//w:pBdr',w)), 0)


if __name__ == '__main__':
    unittest.main(verbosity=2)

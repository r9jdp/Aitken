"""Build the comparison report from saved evidence; never train or change a model.

Run with the bundled document Python. All generated files resolve inside this
comparison directory. Source observations and historic results are read-only.
"""
import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
NAME = 'London_PM25_Comparison_and_Error_Analysis'
BLUE, PALE, GRAY = '233D52', 'F1F5F8', 'D9D9D9'


def output_path(relative):
    path = (HERE / relative).resolve()
    if not path.is_relative_to(HERE):
        raise ValueError('Report outputs must stay inside Aitken/comparison')
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def digest(path):
    # Git normalises these text inputs. Ignore CRLF vs LF so a fresh checkout
    # reproduces the audit, while still detecting changed content.
    data = path.read_bytes()
    if path.suffix in {'.md', '.json', '.csv'}:
        data = data.replace(b'\r\n', b'\n')
    return hashlib.sha256(data).hexdigest()


def hyperlink(paragraph, label, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '145D8D')
    properties.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    properties.append(underline)
    run.append(properties)
    text = OxmlElement('w:t')
    text.set(qn('xml:space'), 'preserve')
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def put_text(paragraph, text, sources):
    for part in re.split(r'(\[\d+\])', text):
        match = re.fullmatch(r'\[(\d+)\]', part)
        if match:
            hyperlink(paragraph, part, sources[int(match.group(1))]['url'])
        elif part:
            paragraph.add_run(part)


def markdown_text(text, sources):
    return re.sub(r'\[(\d+)\]', lambda m: f"[{m.group(0)}]({sources[int(m.group(1))]['url']})", text)


def table(doc, headers, rows, widths, sources):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for col, width in zip(tab.columns, widths):
        col.width = Inches(width)
    repeat = OxmlElement('w:tblHeader')
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, values in enumerate([headers] + rows):
        cells = tab.rows[0].cells if i == 0 else tab.add_row().cells
        cant_split = OxmlElement('w:cantSplit')
        cells[0]._tc.getparent().get_or_add_trPr().append(cant_split)
        for j, value in enumerate(values):
            cell = cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), BLUE if i == 0 else PALE if i % 2 else 'FFFFFF')
            props.append(shade)
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement(f'w:{side}')
                edge.set(qn('w:val'), 'single')
                edge.set(qn('w:sz'), '4')
                edge.set(qn('w:color'), GRAY)
                borders.append(edge)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side, amount in [('top', 85), ('bottom', 85), ('left', 90), ('right', 90)]:
                edge = OxmlElement(f'w:{side}')
                edge.set(qn('w:w'), str(amount))
                edge.set(qn('w:type'), 'dxa')
                margins.append(edge)
            props.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.line_spacing = 1.03
            para.paragraph_format.keep_with_next = False
            put_text(para, str(value), sources)
            for run in para.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def error_figure(evidence):
    """Two quantitative 100-percent bars, not a modified or simulated heatmap."""
    im = Image.new('RGB', (1400, 470), 'white')
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 27)
    bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 29)
    small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 24)
    low_color, high_color = '#456D89', '#AA3438'
    draw.text((30, 20), 'A small group contributes a large share of squared error', font=bold, fill='#172C3D')
    for y, label, low, high in [
        (112, 'Test observations', evidence['below_25']['row_percent'], evidence['at_least_25']['row_percent']),
        (250, 'Squared error', evidence['below_25']['squared_error_percent'], evidence['at_least_25']['squared_error_percent']),
    ]:
        draw.text((30, y - 39), label, font=font, fill='#172C3D')
        x, width = 30, 1330
        split = x + round(width * low / 100)
        draw.rectangle((x, y, split, y + 49), fill=low_color)
        draw.rectangle((split, y, x + width, y + 49), fill=high_color)
        draw.text((x, y + 56), f'Below 25: {low:.2f}%', font=small, fill=low_color)
        label_high = f'At least 25: {high:.2f}%'
        draw.text((x + width - draw.textlength(label_high, font=small), y + 56), label_high, font=small, fill=high_color)
    draw.text((30, 386), 'Original hybrid model | 9,619 unchanged station-days | 2024 benchmark', font=small, fill='#38454E')
    draw.text((30, 422), '25 is a PM2.5 concentration in µg/m³, not an automatic bad-data threshold.', font=small, fill='#38454E')
    path = output_path('figures/error_concentration.png')
    im.save(path)
    return path


def content(e, metric_rows):
    """One shared content tree prevents the Word and GitHub reports diverging."""
    blocks = []
    def add(kind, *args):
        blocks.append((kind, args))
    def p(text): add('p', text)
    def h(text): add('h2', text)
    def page(title): add('page', title)
    def t(headers, rows, widths): add('table', headers, rows, widths)

    add('title', 'London PM25 Model Comparison and High Pollution Error Analysis')
    p('Aitken project • Technical review in simple terms • 12 September 2026')
    h('1 Main conclusion')
    p('Keep the original hybrid HGB plus residual U-Net as the current project model. It has the lowest numerical error in our existing matched comparison, but its improvement over HGB alone is small. The main unresolved weakness is underprediction during high-pollution episodes. The completed training-value capping test did not provide a useful final improvement. [1] [9]')
    h('Our original model comparison')
    t(['Model', 'MAE', 'RMSE', 'R²'], [[r['model'], f"{float(r['mae_ug_m3']):.4f}", f"{float(r['rmse_ug_m3']):.4f}", f"{float(r['r2']):.4f}"] for r in metric_rows], [3.75, 1.08, 1.08, 1.09])
    p('All rows use the same 9,619 LAQN station-days in 2024, after training on 2021–2023. MAE is the average absolute miss; RMSE penalises large misses more strongly. Both are in µg/m³; lower is better. R² is unitless; higher is better. Seeds are different training initialisations whose predictions are averaged. These are the existing release results, not newly retrained models. [1]')
    p('The hybrid reduces RMSE by only 0.0155 µg/m³, about 0.41%, compared with unconstrained HGB. The saved paired calendar-day bootstrap interval for HGB-minus-hybrid RMSE is approximately −0.00003 to +0.0283 µg/m³. Because it includes zero, the numerical lead is not decisive evidence of superiority. [1]')
    p('High PM2.5 readings make up 2.25% of this benchmark but contribute 38.76% of its squared error. This report explains that failure without removing difficult observations or claiming that high values are automatically noise.')

    page('2 What our product does and how it was evaluated')
    h('The model in plain language')
    p('HGB means histogram-based gradient boosting. HGB and XGBoost learn from tables of environmental and location features using boosted decision trees. The ANN is a fully connected artificial neural network. Standalone U-Net learns a pollution map directly from spatial input layers. The hybrid first makes an HGB estimate, then uses U-Net to learn a spatial correction. [1] [8]')
    p('The implementation combines the hybrid correction in standardised logarithmic concentration space, then converts the result back to µg/m³. The correction multiplier is 0.425. Therefore, “HGB estimate plus U-Net correction” is a useful explanation, but it is not a literal addition of two independent raw-concentration maps. The standalone U-Net has 66 input channels and no HGB input; the hybrid has the additional HGB map. [1]')
    h('Data and output')
    p('The output is a daily 1 km Greater London grid in the British National Grid coordinate system. Each tensor contains 48 × 64 cells; 1,719 cells intersect London. Training combines LAQN monitoring observations and Breathe London observations, while the reported benchmark uses LAQN observations. Environmental inputs include ERA5 weather, Sentinel-5P atmospheric products, Sentinel-2 surface indices, land and location features, and earlier PM2.5 observations. The final model excludes five fire and smoke channels. [1]')
    p('The pipeline already filters invalid hourly values, removes duplicate hours, requires at least 18 valid hours for a daily monitor value, and uses medians when combining monitors in a cell. These checks do not prove that every remaining value is error-free, but a reading above 25 µg/m³ is not sufficient evidence to delete it. [1] [9]')
    h('What the benchmark does and does not establish')
    p('The saved predictions cover 349 evaluated dates from 1 January to 14 December 2024, not every day of the calendar year. A station-day means one evaluated station on one date. Missing observations are not filled in to improve the score. The aggregate errors are calculated over station-days, so dates with more reporting stations receive more weight.')
    p('No same-day PM2.5 observation is used as an input. Earlier PM2.5 readings can be available during rolling evaluation. However, same-day environmental data and retrospective products are used, so this is retrospective daily mapping or nowcasting, not a demonstrated operational future forecast. The 2024 benchmark has already been examined and is not an untouched confirmatory test. [1]')
    p('Monitor-based accuracy does not validate every map cell or establish performance at entirely unseen stations. A bright area on a heatmap is a higher model estimate, not proof of a measured hotspot. Changing the colour scale improves readability but does not change prediction accuracy.')

    page('3 Comparison with published research')
    p('The most relevant main comparison is the Greater London daily 1 km ensemble study by Danesh Yazdi et al. [2]. Recent London studies and a recent U-Net study are also included. This is a focused literature comparison, not a systematic review or a reproduction of those studies.')
    t(['Study and task', 'Reported R²', 'Reported RMSE', 'Important difference'], [
        ['Our original hybrid [1]\nDaily London maps', '0.518', '3.777', '9,619 LAQN station-days; 2024 temporal benchmark.'],
        ['Danesh Yazdi et al., 2020 [2]\nDaily 1 km London', '0.828', '4.231', '2005–2013; monitor-held-out 10-fold CV; AOD and augmented target archive.'],
        ['Dimakopoulou et al., 2022 [3]\nDaily London exposure', '0.66–0.83\nacross models', 'Not used here', '2009–2013; combinations include regional pollution and dispersion information.'],
        ['Schneider et al., 2020 [4]\nDaily 1 km Great Britain', '0.767\nmean CV', 'Not verified here', 'Nationwide, not London-only; satellite and atmospheric-model inputs.'],
        ['Legaria-Santiago et al., 2026 [5]\nHourly London sites', '0.53 / 0.69', '3.32 / 2.41', 'Scenario 2: Marylebone / Camden; traffic and neighbouring pollution.'],
        ['Galindo-Prieto et al., 2026 [6]\nDaily London station CT3', '0.796', '2.261\nRMSEP', 'Original-scale PLS test result; other-station information; not a city-wide map.'],
    ], [2.05, .91, .91, 3.13])
    p('CV means cross-validation; AOD means aerosol optical depth; PLS means partial least squares regression. RMSE and RMSEP are in µg/m³. “Not verified” or “not used” does not mean zero. The 0.66–0.83 range is across models, not a confidence interval; that paper’s hybrids report R² 0.81 and 0.79 and are not HGB plus U-Net. [3]')
    p('These are author-reported results, not a head-to-head ranking. The main London study reports R² 0.828 and RMSE 4.231; our hybrid reports 0.518 and 3.777. Years, pollution variability, targets and validation differ. Some papers use regression-based R²; ours compares prediction squared error with observed variation around its mean. A common benchmark would be needed to establish superiority. [1] [2] [4]')
    h('Recent research using a U Net')
    p('AirQ-ResUNet, published in 2025, uses a residual U-Net-style model to emulate an air-pollution dispersion model in Oslo. It supports the relevance of spatial neural networks, but uses a different city and simulator-derived targets. The accessible publisher preview did not provide verifiable numerical performance, so no score is invented here. [7]')

    page('4 High concentrations account for disproportionate error')
    p('The following values were recomputed from the original hybrid’s saved predictions. Every one of the 9,619 evaluation observations was retained. “High” means observed PM2.5 at least 25 µg/m³ for this diagnostic only; it is not AQI, a health classification or a rule for identifying bad measurements.')
    t(['Observed group', 'Count', 'MAE', 'RMSE', 'Bias', 'Squared error share'], [
        ['All observations', f"{e['overall']['n']:,}", f"{e['overall']['mae']:.2f}", f"{e['overall']['rmse']:.2f}", f"{e['overall']['bias']:.2f}", '100.00%'],
        ['Below 25', f"{e['below_25']['n']:,}", f"{e['below_25']['mae']:.2f}", f"{e['below_25']['rmse']:.2f}", f"{e['below_25']['bias']:.2f}", f"{e['below_25']['squared_error_percent']:.2f}%"],
        ['At least 25', f"{e['at_least_25']['n']:,}", f"{e['at_least_25']['mae']:.2f}", f"{e['at_least_25']['rmse']:.2f}", f"{e['at_least_25']['bias']:.2f}", f"{e['at_least_25']['squared_error_percent']:.2f}%"],
    ], [2.0, .75, .90, .90, .90, 1.55])
    p('MAE, RMSE and bias are in µg/m³. Bias is prediction minus observation; a negative value means underprediction. Source: comparison/evidence/high_pollution_diagnosis.json, generated without training or changing observations.')
    add('figure', 'figures/error_concentration.png', 'Figure 1. Concentration of squared error in the high-PM2.5 group. Both bars total 100%.')
    p('All 216 high observations were underpredicted. Their mean observed concentration was 30.42 µg/m³, but the mean prediction was 16.35 µg/m³. The typical absolute miss was 14.07 µg/m³, compared with 2.07 µg/m³ below 25.')
    p('R² and RMSE are strongly affected by large misses because they use squared errors. Missing by 20 contributes 400 squared units; missing by 5 contributes 25. The larger miss counts 16 times as much. This explains why a small number of difficult observations can substantially lower the overall R².')
    p('This is an error-concentration diagnosis, not proof that those measurements are noise. Deleting high observations from the test would change what is being evaluated rather than solve the prediction failure.')

    page('5 Episode examples and likely reasons')
    p('The first five rows below are the dates contributing the most squared error in the saved benchmark. They are retrospective diagnostic examples, not newly selected test cases. The final row is the previously used ordinary-day map example. Means refer only to stations evaluated on that date, not an area-weighted city average. Means and RMSE are in µg/m³.')
    episode_rows = []
    for d in e['highest_error_days'][:5] + [e['ordinary_example']]:
        episode_rows.append([d['date'], str(d['n']), f"{d['observed_mean']:.2f}", f"{d['predicted_mean']:.2f}", f"{d['rmse']:.2f}", f"{d['squared_error_percent']:.2f}%"])
    t(['Date', 'Sites', 'Observed mean', 'Predicted mean', 'RMSE', 'Share of total squared error'], episode_rows, [1.20, .55, 1.30, 1.30, .90, 1.75])
    p('On 11 March, 33 stations averaged 34.76 µg/m³ while the hybrid averaged 11.09. That one day contributed 13.68% of the entire benchmark’s squared error. March contributed 29.94%. The July example’s lower error must not be presented as representative of all pollution conditions.')
    h('A mismatch between the training loss and the score')
    p('The fixed HGB prediction gives 71% weight to a log-target model and 29% to a direct-concentration model. The U-Net correction is also learned in log space with SmoothL1 loss. Log compression and SmoothL1 give very large absolute errors less dominance than raw squared error does. This could help explain conservative peak predictions, but it does not prove causation: the choices helped earlier validation, and the HGB blend still includes a direct model. [1]')
    h('Limited information about an emerging episode')
    p('Our PM2.5 history uses earlier days, not today’s measured concentrations. Some London studies also use same-day regional pollution, aerosol optical depth or atmospheric-model information. Those inputs can convey a regional episode more directly. Our Sentinel-5P aerosol index is not the same measurement as aerosol optical depth. A controlled input comparison is needed to measure the effect. [1] [2] [3] [4]')
    h('What we have not established')
    p('A peak across many stations is consistent with a broader episode, but this analysis has not attributed the March event to weather, transported smoke, chemistry or sensor faults. Possible site-versus-grid differences, source differences and changing conditions also need investigation. None of these explanations justifies labelling every high reading an outlier.')

    page('6 What the completed capping experiment found')
    p('We tested gentle upper winsorisation: values above a training-only percentile were capped, not replaced with the median. Input features, masks, weights, scaling, architecture and loss stayed fixed. Only experimental training-target copies changed; HGB maps and all validation/test observations stayed unchanged. All error measures below are in µg/m³. [9] [10]')
    h('Development comparison on 2023')
    t(['Model and treatment', 'RMSE', 'MAE', 'MAE at least 25', 'Decision'], [
        ['Standalone control', '3.2897', '2.2155', '8.1180', 'Reference'],
        ['Standalone 99.5th cap', '3.2785', '2.2065', '8.0964', 'Passed'],
        ['Standalone 99th cap', '3.2348', '2.1834', '7.6930', 'Selected'],
        ['Hybrid control', '3.1452', '2.1044', '6.9546', 'Reference'],
        ['Hybrid 99.5th cap', '3.1456', '2.1052', '6.9678', 'Failed'],
        ['Hybrid 99th cap', '3.1410', '2.1045', '6.9987', 'Failed'],
    ], [2.45, 1, 1, 1.4, 1.15])
    p('Training used 2021–2022, seed 42 and matched fresh initialisation. The 99.5th and 99th percentile caps were 43.83 and 38.58 µg/m³, changing 551 and 1,101 training cell targets. A cap had to improve RMSE without worsening overall or high-pollution MAE. Neither hybrid treatment passed all conditions. [9]')
    h('Standalone confirmation on the original 2024 observations')
    t(['Matched three-seed model', 'MAE', 'RMSE', 'R²', 'High MAE'], [
        ['Fresh uncapped control', '2.4464', '4.0592', '0.4430', '15.9819'],
        ['99th percentile cap', '2.4728', '4.1435', '0.4196', '16.7066'],
    ], [2.50, 1.05, 1.05, 1.10, 1.30])
    p('After fitting 2021–2023 with seeds 42, 11 and 22, the selected cap was 35.7735 µg/m³ and changed 2,035 training cell-days. RMSE became worse by 0.0843 µg/m³. The paired seven-day-block 95% interval was +0.0016 to +0.1790, from 5,000 resamples. High-pollution MAE also worsened by 0.7247 µg/m³. [9]')
    p('These fresh matched controls are not the historic standalone row on page 1; do not swap them into the product table or compare an experimental cap with an unmatched old run. The interval measures calendar-block variation for the fitted ensembles, not every possible training run. No hybrid cap was retested on 2024 after failing development screening. [9]')
    p('Verdict: do not adopt capping. The original implementation remains the product. Keeping the negative experiment in GitHub records what was tested; it does not activate that treatment or identify faulty measurements.')

    page('7 Limitations and recommended next steps')
    h('What can be claimed now')
    p('The product produces daily London PM2.5 estimates and maps, and supports comparison among several models on a common benchmark. The hybrid has the best numerical release result, with a small and statistically uncertain advantage over HGB alone. Standalone U-Net did not outperform the tree baselines in this experiment. [1]')
    p('The model is not a proven noise detector, an explanation of atmospheric causes or a system validated at every London location. U-Net is not guaranteed to outperform a tree model simply because it is a spatial neural network.')
    h('Priorities for later controlled experiments')
    t(['Priority', 'Proposed check', 'What would count as evidence'], [
        ['1  Audit episodes', 'Trace units, timestamps, monitor quality flags and nearby station agreement for large misses. Preserve high readings unless a concrete quality issue is found.', 'A documented data defect, or retained measurements with a clear audit trail.'],
        ['2  Match the objective', 'Compare the existing loss/blend with carefully chosen raw-error or peak-aware alternatives on development years only.', 'Better overall RMSE without sacrificing ordinary-day or peak MAE, across matched seeds.'],
        ['3  Check residual training', 'Use out-of-fold HGB training predictions for the U-Net correction instead of in-sample baseline maps.', 'A matched experiment tests whether the residual training task becomes more realistic.'],
        ['4  Check missing information', 'Test episode-relevant atmospheric inputs only if reliable data and their time of availability can be established.', 'An ablation isolates the added input; no same-day target leakage.'],
        ['5  Validate the claim', 'Use new time periods and station-held-out tests; evaluate peak errors and spatial behaviour separately.', 'Results support the intended future-date and unseen-location use cases.'],
    ], [1.15, 3.45, 2.40])
    p('These are recommendations, not changes made for this report. The current pipeline already contains source/density weighting and earlier tuning; proposed alternatives need to be compared against those existing choices. Do not keep tuning against the already examined 2024 benchmark. [1] [9]')
    h('Suggested wording for faculty')
    p('“Our hybrid combines an HGB estimate with a U-Net spatial correction. It is the best numerical model in our current comparison, with RMSE 3.78 µg/m³ and R² 0.518, although the gain over HGB is small. Our main weakness is pollution peaks: 216 high observations cause about 39% of squared error. We tested gentle training-data capping, but it did not improve the final result, so we retained the original model. Published London studies are useful reference points, but their data and evaluation settings differ.”')
    h('Reproducibility record')
    p('Own-model scores come from results/metrics.csv; the completed cap test comes from results/outlier_sensitivity_20260909. New aggregates and the source-prediction SHA-256 are in comparison/evidence/high_pollution_diagnosis.json. comparison/README.md gives reproduction commands, and sources.json records paper tables and caveats. This report adds no training, observations or model changes.')

    page('8 References')
    p('Numbered references are clickable in both report formats. Publication details and quantitative claims were checked against primary publisher pages or author manuscripts on 12 September 2026. Older London papers are retained for direct task relevance; the 2025–2026 studies provide recent context.')
    add('references')
    return blocks


def build():
    inputs = [ROOT/'results/metrics.csv', ROOT/'docs/RESEARCH_PROTOCOL.md',
              ROOT/'results/outlier_sensitivity_20260909/results.json',
              ROOT/'results/outlier_sensitivity_20260909/audit.json',
              HERE/'evidence/high_pollution_diagnosis.json', HERE/'sources.json']
    before = {p.relative_to(ROOT).as_posix(): digest(p) for p in inputs}
    sources = {s['id']: s for s in json.loads((HERE/'sources.json').read_text(encoding='utf-8'))['sources']}
    e = json.loads((HERE/'evidence/high_pollution_diagnosis.json').read_text(encoding='utf-8'))
    with (ROOT/'results/metrics.csv').open(newline='', encoding='utf-8') as stream:
        metrics = list(csv.DictReader(stream))
    assert len(metrics) == 7 and e['at_least_25']['underprediction_count'] == 216
    assert e['verification']['published_hybrid_metrics_match']
    error_figure(e)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.70), Inches(.65)
    section.left_margin, section.right_margin = Inches(.75), Inches(.75)
    section.header_distance, section.footer_distance = Inches(.25), Inches(.25)
    for name, size in [('Normal', 11), ('Title', 22), ('Heading 1', 17), ('Heading 2', 12), ('Caption', 9)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7 if name == 'Normal' else 8)
        style.paragraph_format.line_spacing = 1.06
    doc.styles['Normal'].paragraph_format.widow_control = True
    # The bundled python-docx base template includes a decorative Title rule.
    # This report intentionally uses plain, black headings with no title rule.
    for style in doc.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Title'].font.bold = True
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
    header = section.header.paragraphs[0]
    header.add_run('AITKEN  /  LONDON PM2.5  /  MODEL REVIEW').font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run('12 September 2026  •  Retrospective benchmark').font.size = Pt(8)
    footer.add_run(' '*8 + 'Page ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = 'London PM25 Model Comparison and High Pollution Error Analysis'
    doc.core_properties.subject = 'Original results, published research and high-concentration prediction failures'
    doc.core_properties.author = 'Aitken project'
    doc.core_properties.keywords = 'London, PM2.5, U-Net, HGB, evaluation, error analysis'
    md = []
    blocks = content(e, metrics)
    for kind, args in blocks:
        if kind in ('title', 'page', 'h2'):
            if kind == 'page':
                doc.add_page_break()
            text = args[0]
            doc.add_paragraph(text, style={'title':'Title','page':'Heading 1','h2':'Heading 2'}[kind])
            md.extend([{'title':'# ','page':'## ','h2':'### '}[kind] + text, ''])
        elif kind == 'p':
            para = doc.add_paragraph()
            put_text(para, args[0], sources)
            md.extend([markdown_text(args[0], sources), ''])
        elif kind == 'table':
            headers, rows, widths = args
            assert abs(sum(widths)-7) < 1e-7
            table(doc, headers, rows, widths, sources)
            for row in [headers, ['---']*len(headers)] + rows:
                md.append('| ' + ' | '.join(markdown_text(str(v).replace('\n','<br>'), sources) for v in row) + ' |')
            md.append('')
        elif kind == 'figure':
            path, caption = args
            doc.add_picture(str(HERE/path), width=Inches(6.95))
            doc.add_paragraph(caption, style='Caption')
            md.extend([f'![{caption}]({path})', ''])
        elif kind == 'references':
            for n, source in sources.items():
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(9)
                para.paragraph_format.line_spacing = 1.02
                hyperlink(para, f'[{n}] ', source['url'])
                run = para.add_run(source['reference'])
                run.font.size = Pt(10)
                para.add_run(' ')
                hyperlink(para, 'Online source', source['url'])
                md.extend([f"[{n}] {source['reference']} [Online source]({source['url']}).", ''])
    doc.save(output_path(NAME + '.docx'))
    output_path('REPORT.md').write_text('\n'.join(md), encoding='utf-8', newline='\n')
    assert before == {p.relative_to(ROOT).as_posix(): digest(p) for p in inputs}
    audit = dict(status='passed', input_hashes_unchanged=before,
                 input_hash_method='SHA-256 with CRLF normalised to LF for tracked text inputs; original prediction gzip uses byte-exact SHA-256 in the diagnosis',
                 generated=[NAME+'.docx','REPORT.md','figures/error_concentration.png'],
                 training_performed=False, observations_modified=False,
                 dashboard_presentation_and_product_table_modified=False,
                 output_scope='Aitken/comparison only', primary_references=len(sources),
                 note='Rendering and visual review are recorded separately in REVIEW.md.')
    output_path('evidence/build_audit.json').write_text(json.dumps(audit, indent=2), encoding='utf-8', newline='\n')
    print(json.dumps(audit, indent=2))


if __name__ == '__main__':
    build()
